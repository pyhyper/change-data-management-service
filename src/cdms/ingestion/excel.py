import io
from typing import Any
import openpyxl
from ..domain.exceptions import BadRequestException
from ..domain.models import NormalizedChange
from ..processing.canonical import compute_payload_hash, normalize_product
from ..processing.idempotency import compute_buffer_hash, derive_excel_idempotency_key
from ..processing.pipeline import ChangeProcessor

REQUIRED_HEADERS = {"id", "updatedat"}

def parse_excel_buffer(buffer: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(buffer), data_only=True)
    except Exception as e:
        raise BadRequestException(f"Failed to parse Excel file: {str(e)}")

    sheet_names = wb.sheetnames
    if not sheet_names:
        raise BadRequestException("The uploaded Excel workbook contains no sheets.")

    sheet = wb[sheet_names[0]]
    rows_iter = sheet.iter_rows(values_only=True)

    header_row = next(rows_iter, None)
    if not header_row or all(c is None for c in header_row):
        raise BadRequestException("The uploaded Excel sheet contains no header row.")

    headers = [str(c).strip() if c is not None else "" for c in header_row]
    lower_headers = [h.lower() for h in headers]

    for req in REQUIRED_HEADERS:
        if req not in lower_headers:
            raise BadRequestException(
                f"Missing required header '{req}' in Excel sheet. Found: {', '.join(headers)}"
            )

    header_map = {}
    for idx, h_low in enumerate(lower_headers):
        if h_low == "updatedat":
            header_map[idx] = "updatedAt"
        else:
            header_map[idx] = h_low

    data_rows = []
    for r in rows_iter:
        if all(c is None or str(c).strip() == "" for c in r):
            continue
        row_dict = {}
        for idx, col_name in header_map.items():
            if idx < len(r):
                row_dict[col_name] = r[idx]
        data_rows.append(row_dict)

    return headers, data_rows

async def process_excel_upload(
    file_bytes: bytes,
    filename: str,
    change_processor: ChangeProcessor
) -> dict[str, Any]:
    file_hash = compute_buffer_hash(file_bytes)
    _, rows = parse_excel_buffer(file_bytes)

    summary = {
        "file": filename,
        "rows": len(rows),
        "inserted": 0,
        "noChange": 0,
        "stale": 0,
        "duplicate": 0,
        "invalid": 0,
        "errors": [],
    }

    # Process rows with batching
    batch_size = 10
    for i in range(0, len(rows), batch_size):
        batch = rows[i:i + batch_size]
        for batch_idx, row in enumerate(batch):
            row_num = i + batch_idx + 2
            try:
                normalized = normalize_product(row)
                payload_hash = compute_payload_hash(normalized)
                idempotency_key = derive_excel_idempotency_key(
                    file_hash,
                    row_num,
                    normalized["id"],
                    payload_hash
                )
                from datetime import datetime
                change = NormalizedChange(
                    product_id=normalized["id"],
                    source="EXCEL",
                    source_updated_at=datetime.fromisoformat(normalized["updatedAt"].replace("Z", "+00:00")),
                    payload=normalized,
                    payload_hash=payload_hash,
                    idempotency_key=idempotency_key,
                )
                res = await change_processor.process_change(change)
                if res.status == "INSERTED":
                    summary["inserted"] += 1
                elif res.status == "NO_CHANGE":
                    summary["noChange"] += 1
                elif res.status == "STALE":
                    summary["stale"] += 1
                elif res.status == "DUPLICATE_EVENT":
                    summary["duplicate"] += 1
            except Exception as e:
                summary["invalid"] += 1
                summary["errors"].append({"row": row_num, "error": str(e)})

    return summary
